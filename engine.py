"""Deterministic purchasing engine for SmartZakup.

The module deliberately has no Streamlit imports.  All quantities are calculated
from the supplied Excel files and the formulas in SPEC.md.
"""

from __future__ import annotations

import calendar
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd


SUPPLIERS = ["Systeme Electric", "IEK"]


@dataclass
class Settings:
    today: date = date(2026, 9, 22)
    demand_source: str = "invoices"
    lead_time_days: dict = field(
        default_factory=lambda: {"IEK": 35, "Systeme Electric": 30}
    )
    review_days: dict = field(
        default_factory=lambda: {"IEK": 14, "Systeme Electric": 30}
    )
    service_z: dict = field(
        default_factory=lambda: {"A": 1.65, "B": 1.28, "C": 0.84}
    )
    plan_growth_pct: float = 0.0


PLAN_COLUMNS = [
    "supplier", "code", "article", "name", "category", "abc_class", "unit",
    "available", "in_transit", "next_arrival", "base_monthly", "trend",
    "window_days", "window_demand", "safety_stock", "raw_qty", "multiple",
    "qty", "urgency", "enough_until", "price", "value", "do_not_reorder",
    "reason",
]


MONTH_PREFIXES = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6,
    "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12,
}


def _read_excel(path: Path, **kwargs) -> pd.DataFrame:
    """Use the fast reader and fall back for files with unusual Excel features."""
    try:
        return pd.read_excel(path, engine="calamine", **kwargs)
    except Exception:
        return pd.read_excel(path, engine="openpyxl", **kwargs)


def _text(value) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def _number(series: pd.Series) -> pd.Series:
    return pd.to_numeric(
        series.astype(str).str.replace("\u00a0", "", regex=False)
        .str.replace(" ", "", regex=False).str.replace(",", ".", regex=False),
        errors="coerce",
    ).fillna(0.0)


def _month_from_label(label) -> pd.Period | None:
    text = _text(label).lower().replace("ё", "е")
    year_match = re.search(r"(20\d{2})", text)
    if not year_match:
        return None
    for prefix, month in MONTH_PREFIXES.items():
        if prefix in text:
            return pd.Period(year=int(year_match.group(1)), month=month, freq="M")
    return None


def _cached_table(
    cache_dir: Path, name: str, sources: list[Path], builder: Callable[[], pd.DataFrame]
) -> pd.DataFrame:
    cache_path = cache_dir / f"{name}.parquet"
    newest_source = max(path.stat().st_mtime for path in sources)
    if cache_path.exists() and cache_path.stat().st_mtime >= newest_source:
        return pd.read_parquet(cache_path)
    result = builder()
    cache_dir.mkdir(parents=True, exist_ok=True)
    result.to_parquet(cache_path, index=False)
    return result


def _load_invoices(data_dir: Path) -> pd.DataFrame:
    frames = []
    for supplier, filename in [
        ("Systeme Electric", "se_sales_invoices.xlsx"),
        ("IEK", "iek_sales_invoices.xlsx"),
    ]:
        raw = _read_excel(data_dir / filename, header=0, dtype=str)
        raw = raw.iloc[:, :8].copy()
        raw.columns = ["date", "doc_number", "document", "code", "name", "unit", "warehouse", "qty"]
        raw["date"] = pd.to_datetime(raw["date"], dayfirst=True, errors="coerce")
        raw["document"] = raw["document"].map(_text)
        raw["code"] = raw["code"].map(_text)
        raw["qty"] = _number(raw["qty"]).abs()
        keep = (
            raw["date"].notna()
            & (raw["date"] >= pd.Timestamp("2025-01-01"))
            & raw["document"].str.startswith("Расходная накладная", na=False)
            & raw["code"].ne("")
            & raw["qty"].gt(0)
        )
        clean = raw.loc[keep, ["date", "doc_number", "code", "name", "unit", "qty"]].copy()
        clean.rename(columns={"doc_number": "doc"}, inplace=True)
        for col in ["doc", "name", "unit"]:
            clean[col] = clean[col].map(_text)
        clean.insert(0, "supplier", supplier)
        frames.append(clean)
    return pd.concat(frames, ignore_index=True)


def _load_monthly(data_dir: Path, kind: str) -> pd.DataFrame:
    frames = []
    files = {
        "sales": [("Systeme Electric", "se_sales_monthly.xlsx", 1), ("IEK", "iek_sales_monthly.xlsx", 1)],
        "stock": [("Systeme Electric", "se_stock_monthly.xlsx", 2), ("IEK", "iek_stock_monthly.xlsx", 2)],
    }
    for supplier, filename, code_pos in files[kind]:
        raw = _read_excel(data_dir / filename, header=0, dtype=str)
        code_col = raw.columns[code_pos]
        month_cols = [col for col in raw.columns if _month_from_label(col) is not None]
        slim = raw[[code_col, *month_cols]].copy()
        slim.rename(columns={code_col: "code"}, inplace=True)
        slim["code"] = slim["code"].map(_text)
        slim = slim[slim["code"].ne("") & slim["code"].ne("Итого")]
        long = slim.melt(id_vars="code", var_name="month_label", value_name="value")
        long["month"] = long["month_label"].map(_month_from_label)
        long["value"] = _number(long["value"])
        if kind == "sales":
            long["value"] = long["value"].clip(lower=0)
            value_name = "qty"
        else:
            value_name = "stock_start"
        long.rename(columns={"value": value_name}, inplace=True)
        long.insert(0, "supplier", supplier)
        frames.append(long[["supplier", "code", "month", value_name]])
    return pd.concat(frames, ignore_index=True)


def _series_from_name(name: str) -> str:
    matches = re.findall(r'["«](.*?)["»]', name)
    return matches[-1].strip() if matches else "Прочее"


def _load_catalog_and_manager(data_dir: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    manager = _read_excel(data_dir / "se_manager_sheet.xlsx", sheet_name="TDSheet", header=1, dtype=str)
    manager.columns = [re.sub(r"\s+", " ", _text(c)) for c in manager.columns]
    manager = manager.rename(columns={
        "Код 1с": "code", "Артикул поставщика": "article", "Наименование": "name",
        "Категория 2026": "manager_class", "СС реал": "price",
        "Витрина": "showcase", "Остаток ТЗ": "tz_stock",
        "Розничный склад": "retail_stock", "Свободный остаток": "free_stock",
        "СЭ в пути 24.09": "transit_qty",
    })
    for col in ["code", "article", "name", "manager_class"]:
        manager[col] = manager[col].map(_text)
    manager = manager[manager["code"].ne("")].copy()
    for col in ["price", "showcase", "tz_stock", "retail_stock", "free_stock", "transit_qty"]:
        manager[col] = _number(manager[col]) if col in manager else 0.0
    manager["available"] = manager[["showcase", "tz_stock", "retail_stock", "free_stock"]].sum(axis=1)
    manager["arrival_date"] = date(2026, 9, 24)

    # Arrival dates sometimes live in comments on column BC.
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(data_dir / "se_manager_sheet.xlsx", read_only=False, data_only=True)
        sheet = workbook["TDSheet"]
        comment_dates = {}
        for row in range(3, sheet.max_row + 1):
            code = _text(sheet.cell(row=row, column=3).value)
            comment = sheet.cell(row=row, column=55).comment
            if not code or not comment:
                continue
            match = re.search(r"(\d{1,2})[.](\d{1,2})(?:[.](\d{2,4}))?", comment.text)
            if match:
                year = int(match.group(3) or 2026)
                if year < 100:
                    year += 2000
                comment_dates[code] = date(year, int(match.group(2)), int(match.group(1)))
        manager["arrival_date"] = manager.apply(
            lambda row: comment_dates.get(row["code"], row["arrival_date"]), axis=1
        )
    except Exception:
        pass

    se_catalog = pd.DataFrame({
        "supplier": "Systeme Electric",
        "code": manager["code"],
        "article": manager["article"],
        "name": manager["name"],
        "category": manager["name"].map(_series_from_name),
        "unit": "",
        "status": "",
        "price": manager["price"],
        "manager_class": manager["manager_class"],
        "catalog_multiple": 1,
    })

    iek = _read_excel(data_dir / "iek_catalog.xlsx", header=0, dtype=str)
    iek.columns = [re.sub(r"\s+", " ", _text(c)) for c in iek.columns]
    iek = iek.rename(columns={
        "Код 1с": "code", "Артикул": "article", "Наименование": "name",
        "Категория": "category", "Ед.": "unit", "Статус IEK": "status",
        "Цена IEK, тенге с НДС": "price", "Мин. отгрузка": "catalog_multiple",
    })
    for col in ["code", "article", "name", "category", "unit", "status"]:
        iek[col] = iek[col].map(_text)
    iek = iek[iek["code"].ne("")].copy()
    iek["price"] = _number(iek["price"])
    iek["catalog_multiple"] = _number(iek["catalog_multiple"]).clip(lower=1).map(math.ceil)
    iek["manager_class"] = ""
    iek.insert(0, "supplier", "IEK")
    iek = iek[["supplier", "code", "article", "name", "category", "unit", "status", "price", "manager_class", "catalog_multiple"]]
    catalog = pd.concat([se_catalog, iek], ignore_index=True).drop_duplicates(["supplier", "code"], keep="first")
    return catalog, manager[["code", "manager_class", "available", "transit_qty", "arrival_date"]]


def _load_transit(data_dir: Path, manager: pd.DataFrame) -> pd.DataFrame:
    parts = []
    se = manager[manager["transit_qty"].gt(0)].rename(columns={"transit_qty": "qty"})
    if not se.empty:
        se = se[["code", "qty", "arrival_date"]].copy()
        se.insert(0, "supplier", "Systeme Electric")
        parts.append(se)

    raw = _read_excel(data_dir / "iek_in_transit.xlsx", sheet_name="Лист4", header=0, dtype=str)
    code_col = raw.columns[0]
    for col in raw.columns[3:]:
        match = re.search(r"(\d{2})[.](\d{2})[.](\d{4})", _text(col))
        if not match:
            continue
        arrival = date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
        frame = pd.DataFrame({"code": raw[code_col].map(_text), "qty": _number(raw[col])})
        frame = frame[frame["code"].ne("") & frame["qty"].gt(0)]
        frame.insert(0, "supplier", "IEK")
        frame["arrival_date"] = arrival
        parts.append(frame)
    if not parts:
        return pd.DataFrame(columns=["supplier", "code", "qty", "arrival_date"])
    result = pd.concat(parts, ignore_index=True)
    return result.groupby(["supplier", "code", "arrival_date"], as_index=False)["qty"].sum()


def _load_moq(data_dir: Path, catalog: pd.DataFrame) -> pd.DataFrame:
    parts = []
    specs = [
        ("Systeme Electric", "se_moq.xlsx", 2, 4),
        ("IEK", "iek_moq.xlsx", 1, 4),
    ]
    for supplier, filename, code_pos, multiple_pos in specs:
        raw = _read_excel(data_dir / filename, header=0, dtype=str)
        frame = pd.DataFrame({
            "supplier": supplier,
            "code": raw.iloc[:, code_pos].map(_text),
            "multiple": _number(raw.iloc[:, multiple_pos]).clip(lower=1).map(math.ceil),
        })
        parts.append(frame[frame["code"].ne("")])
    result = pd.concat(parts, ignore_index=True).drop_duplicates(["supplier", "code"])
    fallback = catalog[["supplier", "code", "catalog_multiple"]].rename(columns={"catalog_multiple": "fallback"})
    # Keep MOQ rows for products absent from the manager/catalog as well.
    result = fallback.merge(result, on=["supplier", "code"], how="outer")
    result["multiple"] = result["multiple"].fillna(result["fallback"]).fillna(1).clip(lower=1).astype(int)
    return result[["supplier", "code", "multiple"]]


def _load_company_seasonality(data_dir: Path) -> pd.DataFrame:
    rows = []
    for supplier, filename in [("Systeme Electric", "se_seasonality.xlsx"), ("IEK", "iek_seasonality.xlsx")]:
        raw = _read_excel(data_dir / filename, header=2, dtype=str)
        year_col = raw.columns[0]
        for year in [2024, 2025]:
            selected = raw[raw[year_col].map(_text).eq(str(year))]
            if selected.empty:
                continue
            values = _number(selected.iloc[0, 1:13])
            mean = values.mean()
            factors = values / mean if mean > 0 else pd.Series([1.0] * 12)
            rows.extend({"supplier": supplier, "year": year, "month_num": i + 1, "factor": float(value)} for i, value in enumerate(factors))
    return pd.DataFrame(rows)


def load_all(data_dir: str = "data") -> dict:
    """Read and clean every source table described by the public interface."""
    root = Path(data_dir)
    cache = root / "cache"
    invoice_sources = [root / "se_sales_invoices.xlsx", root / "iek_sales_invoices.xlsx"]
    monthly_sources = [root / "se_sales_monthly.xlsx", root / "iek_sales_monthly.xlsx"]
    stock_sources = [root / "se_stock_monthly.xlsx", root / "iek_stock_monthly.xlsx"]
    invoices = _cached_table(cache, "invoices", invoice_sources, lambda: _load_invoices(root))
    sales_monthly = _cached_table(cache, "sales_monthly", monthly_sources, lambda: _load_monthly(root, "sales"))
    stock_monthly = _cached_table(cache, "stock_monthly", stock_sources, lambda: _load_monthly(root, "stock"))

    catalog, manager = _load_catalog_and_manager(root)
    invoice_units = invoices.groupby(["supplier", "code"])["unit"].first().to_dict()
    catalog["unit"] = [
        _text(row.unit) or invoice_units.get((row.supplier, row.code), "шт")
        for row in catalog.itertuples()
    ]
    in_transit = _load_transit(root, manager)
    moq = _load_moq(root, catalog)
    company_seasonality = _load_company_seasonality(root)
    return {
        "invoices": invoices,
        "sales_monthly": sales_monthly,
        "stock_monthly": stock_monthly,
        "in_transit": in_transit,
        "moq": moq,
        "catalog": catalog.drop(columns=["manager_class", "catalog_multiple"], errors="ignore"),
        "manager": manager,
        "catalog_meta": catalog,
        "company_seasonality": company_seasonality,
    }


def _normalise_factors(values) -> np.ndarray:
    array = np.asarray(values, dtype=float)
    array = np.where(np.isfinite(array), array, 1.0)
    if array.size != 12 or array.mean() <= 0:
        array = np.ones(12)
    array = np.clip(array, 0.3, 3.0)
    return array / array.mean()


def _shape(values: pd.Series) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    mean = values.mean()
    return values / mean if mean > 0 else np.ones(12)


def _trim(values: pd.Series) -> np.ndarray:
    array = np.asarray(pd.Series(values).dropna(), dtype=float)
    if array.size >= 4:
        # Sorting also removes TWO months when minimum and maximum are tied.
        return np.sort(array)[1:-1]
    return array


def _window_forecast(
    start: date, days: int, base: float, factors: np.ndarray, trend: float, growth: float
) -> float:
    total = 0.0
    middle = pd.Timestamp("2026-06-01")
    for offset in range(days):
        current = start + timedelta(days=offset)
        months = (current.year - middle.year) * 12 + current.month - middle.month
        monthly = base * factors[current.month - 1] * (1 + trend) ** (months / 12) * (1 + growth / 100)
        total += monthly / calendar.monthrange(current.year, current.month)[1]
    return float(max(0.0, total))


def _reason_date(value) -> str:
    return value.strftime("%d.%m.%Y") if value else "—"


def build_plan(data: dict, settings: Settings, overrides: dict | None = None) -> tuple[pd.DataFrame, dict]:
    """Build one deterministic recommendation per supplier/product."""
    overrides = overrides or {}
    invoices = data["invoices"].copy()
    invoices["date"] = pd.to_datetime(invoices["date"])
    catalog = data.get("catalog_meta", data["catalog"]).copy()
    if "manager_class" not in catalog:
        catalog["manager_class"] = ""
    if "catalog_multiple" not in catalog:
        catalog["catalog_multiple"] = 1

    extras = overrides.get("extra_invoice_lines", [])
    if extras:
        lookup = catalog.set_index(["supplier", "code"])
        extra_rows = []
        for item in extras:
            supplier = item.get("supplier", "IEK")
            code = _text(item.get("code"))
            try:
                product = lookup.loc[(supplier, code)]
                name, unit = product["name"], product["unit"]
            except KeyError:
                name, unit = code, "шт"
            extra_rows.append({
                "supplier": supplier, "date": pd.to_datetime(item["date"]), "doc": "Проверка",
                "code": code, "name": name, "unit": unit, "qty": float(item["qty"]),
            })
        invoices = pd.concat([invoices, pd.DataFrame(extra_rows)], ignore_index=True)

    invoices["month"] = invoices["date"].dt.to_period("M")
    invoice_stats = invoices.groupby(["supplier", "code"])["qty"].agg(line_count="size", typical_qty="median")
    monthly_raw = invoices.groupby(["supplier", "code", "month"], as_index=False)["qty"].sum()
    monthly_median = monthly_raw.groupby(["supplier", "code"])["qty"].median().rename("typical_month")
    invoices = invoices.join(invoice_stats, on=["supplier", "code"]).join(monthly_median, on=["supplier", "code"])
    invoices["is_one_off"] = (
        invoices["line_count"].ge(10)
        & invoices["qty"].gt(20 * invoices["typical_qty"])
        & invoices["qty"].gt(0.5 * invoices["typical_month"])
    )
    if overrides.get("disable_one_off_filter", False):
        invoices["is_one_off"] = False
    regular_invoices = invoices[~invoices["is_one_off"]].copy()
    september_regular = regular_invoices[
        (regular_invoices["date"].dt.date >= date(2026, 9, 1))
        & (regular_invoices["date"].dt.date <= settings.today)
    ]
    september_sales = september_regular.groupby(["supplier", "code"])["qty"].sum().to_dict()
    one_off_frames = {
        key: frame for key, frame in invoices[invoices["is_one_off"]].groupby(["supplier", "code"])
    }
    invoice_clean_monthly = regular_invoices.groupby(["supplier", "code", "month"])["qty"].sum().to_dict()
    invoice_raw_monthly = invoices.groupby(["supplier", "code", "month"])["qty"].sum().to_dict()

    sales = data["sales_monthly"].copy()
    sales["month"] = sales["month"].map(lambda x: pd.Period(x, freq="M"))
    report_monthly = sales.groupby(["supplier", "code", "month"])["qty"].sum().to_dict()
    stock = data["stock_monthly"].copy()
    stock["month"] = stock["month"].map(lambda x: pd.Period(x, freq="M"))
    stock_lookup = stock.groupby(["supplier", "code", "month"])["stock_start"].sum().to_dict()

    # Add products present in operational tables but absent from a catalog.
    known = catalog[["supplier", "code"]]
    extra_products = sales[["supplier", "code"]].drop_duplicates().merge(known, how="left", indicator=True)
    extra_products = extra_products[extra_products["_merge"].eq("left_only")][["supplier", "code"]]
    if not extra_products.empty:
        names = invoices.groupby(["supplier", "code"])[["name", "unit"]].first().reset_index()
        extra_products = extra_products.merge(names, on=["supplier", "code"], how="left")
        extra_products["article"] = ""
        extra_products["category"] = "Прочее"
        extra_products["status"] = ""
        extra_products["price"] = 0.0
        extra_products["manager_class"] = ""
        extra_products["catalog_multiple"] = 1
        catalog = pd.concat([catalog, extra_products[catalog.columns]], ignore_index=True)

    catalog = catalog.drop_duplicates(["supplier", "code"]).copy()
    product_keys = list(catalog[["supplier", "code"]].itertuples(index=False, name=None))
    full_months = pd.period_range("2024-01", "2026-08", freq="M")
    recent12 = pd.period_range("2025-09", "2026-08", freq="M")
    base_months = pd.period_range("2026-03", "2026-08", freq="M")
    prior_months = pd.period_range("2025-03", "2025-08", freq="M")
    season_months = {year: pd.period_range(f"{year}-01", f"{year}-12", freq="M") for year in (2024, 2025)}

    def demand_value(supplier, code, month, cleaned=True):
        if month.year == 2024 or settings.demand_source == "monthly_report":
            return float(report_monthly.get((supplier, code, month), 0.0))
        table = invoice_clean_monthly if cleaned else invoice_raw_monthly
        return float(table.get((supplier, code, month), 0.0))

    # Supplier seasonality: company unit shapes blended 50/50 with the supplied money shape.
    supplier_factors = {}
    company = data.get("company_seasonality", pd.DataFrame())
    for supplier in SUPPLIERS:
        yearly = []
        supplier_codes = catalog.loc[catalog["supplier"].eq(supplier), "code"]
        for year in [2024, 2025]:
            values = [sum(demand_value(supplier, code, month) for code in supplier_codes) for month in season_months[year]]
            yearly.append(_shape(values))
        units_shape = np.mean(yearly, axis=0)
        money_rows = company[company["supplier"].eq(supplier)] if not company.empty else pd.DataFrame()
        money_shape = money_rows.groupby("month_num")["factor"].mean().reindex(range(1, 13), fill_value=1).to_numpy() if not money_rows.empty else np.ones(12)
        supplier_factors[supplier] = _normalise_factors(0.5 * units_shape + 0.5 * money_shape)

    category_factors = {}
    category_counts = catalog.groupby(["supplier", "category"])["code"].nunique()
    for (supplier, category), count in category_counts.items():
        if count < 10:
            continue
        codes = catalog.loc[(catalog["supplier"].eq(supplier)) & (catalog["category"].eq(category)), "code"]
        yearly = []
        for year in [2024, 2025]:
            values = [sum(demand_value(supplier, code, month) for code in codes) for month in season_months[year]]
            yearly.append(_shape(values))
        category_factors[(supplier, category)] = _normalise_factors(np.mean(yearly, axis=0))

    transit = data["in_transit"].copy()
    transit_totals = transit.groupby(["supplier", "code"])["qty"].sum().to_dict()
    arrival_lookup = transit.groupby(["supplier", "code"])["arrival_date"].min().to_dict()
    moq = data["moq"].set_index(["supplier", "code"])["multiple"].to_dict()
    manager = data.get("manager", pd.DataFrame()).set_index("code") if not data.get("manager", pd.DataFrame()).empty else pd.DataFrame()

    # IEK ABC classification by regular last-12-month sales value.
    iek_rows = catalog[catalog["supplier"].eq("IEK")].copy()
    iek_rows["last12"] = [sum(demand_value("IEK", code, month) for month in recent12) for code in iek_rows["code"]]
    iek_rows["importance"] = iek_rows["last12"] * iek_rows["price"].where(iek_rows["price"].gt(0), 1.0)
    iek_rows = iek_rows.sort_values("importance", ascending=False)
    total_importance = iek_rows["importance"].sum()
    cumulative_before = iek_rows["importance"].cumsum() - iek_rows["importance"]
    shares = cumulative_before / total_importance if total_importance > 0 else pd.Series(1.0, index=iek_rows.index)
    iek_rows["abc"] = np.where(shares < 0.80, "A", np.where(shares < 0.95, "B", "C"))
    iek_abc = iek_rows.set_index("code")["abc"].to_dict()

    plan_rows = []
    details = {}
    for product in catalog.itertuples(index=False):
        supplier, code = product.supplier, product.code
        category = _text(product.category) or "Прочее"
        raw_values = pd.Series(
            [demand_value(supplier, code, month, cleaned=False) for month in full_months], index=full_months, dtype=float
        )
        cleaned_values = pd.Series(
            [demand_value(supplier, code, month, cleaned=True) for month in full_months], index=full_months, dtype=float
        )

        category_shape = category_factors.get((supplier, category), supplier_factors[supplier])
        own_2024 = cleaned_values[cleaned_values.index.year == 2024].to_numpy()
        own_2025 = cleaned_values[cleaned_values.index.year == 2025].to_numpy()
        correlation = np.corrcoef(own_2024, own_2025)[0, 1] if np.std(own_2024) > 0 and np.std(own_2025) > 0 else np.nan
        if own_2024.sum() >= 60 and own_2025.sum() >= 60 and np.isfinite(correlation) and correlation >= 0.5:
            own_shape = _normalise_factors((_shape(own_2024) + _shape(own_2025)) / 2)
            factors = _normalise_factors(0.6 * own_shape + 0.4 * category_shape)
            season_level = "product"
        elif (supplier, category) in category_factors:
            factors = category_shape
            season_level = "category"
        else:
            factors = supplier_factors[supplier]
            season_level = "supplier"

        stock_values = pd.Series(
            [float(stock_lookup.get((supplier, code, month), 0.0)) for month in full_months], index=full_months
        )
        last12_clean = cleaned_values.reindex(recent12, fill_value=0.0)
        eligible = [m for m in recent12 if stock_values.get(m, 0) > 0]
        base_pool = eligible if len(eligible) >= 3 else list(recent12)
        deseasoned_pool = [cleaned_values.get(m, 0.0) / factors[m.month - 1] for m in base_pool]
        preliminary_base = float(np.median(deseasoned_pool)) if deseasoned_pool else 0.0
        adjusted_values = cleaned_values.copy()
        stockout_rows = []
        for month in recent12:
            expected = preliminary_base * factors[month.month - 1]
            actual = cleaned_values.get(month, 0.0)
            next_month = month + 1
            start_stock = stock_values.get(month, 0.0)
            next_stock = float(stock_lookup.get((supplier, code, next_month), 0.0))
            is_stockout = (start_stock <= 0 or next_stock <= 0) and actual < 0.7 * expected
            if is_stockout and expected > actual:
                stockout_rows.append({"month": month, "actual": actual, "expected": expected, "lost": expected - actual})
                if not overrides.get("disable_stockout_adjustment", False):
                    adjusted_values.loc[month] = expected

        deseasoned = pd.Series(
            [adjusted_values.get(month, 0.0) / factors[month.month - 1] for month in full_months], index=full_months
        )
        base_sample = deseasoned.reindex(base_months, fill_value=0.0)
        base_array = _trim(base_sample)
        base = float(base_array.mean()) if base_array.size else 0.0
        current_sum = float(_trim(deseasoned.reindex(base_months, fill_value=0.0)).sum())
        prior_sum = float(_trim(deseasoned.reindex(prior_months, fill_value=0.0)).sum())
        trend = float(np.clip(current_sum / prior_sum - 1, -0.4, 0.6)) if current_sum >= 30 and prior_sum >= 30 else 0.0
        variability_values = deseasoned.reindex(recent12, fill_value=0.0).to_numpy()
        median = np.median(variability_values)
        sigma = max(1.4826 * float(np.median(np.abs(variability_values - median))), 0.2 * base)

        if supplier == "Systeme Electric" and not manager.empty and code in manager.index:
            selected = manager.loc[code]
            if isinstance(selected, pd.DataFrame):
                selected = selected.iloc[0]
            available = float(selected["available"])
        else:
            sep_stock = float(stock_lookup.get((supplier, code, pd.Period("2026-09", freq="M")), 0.0))
            sep_sales = september_sales.get((supplier, code), 0.0)
            available = max(0.0, sep_stock - float(sep_sales))
        available = float(overrides.get("stock", {}).get(code, available))
        in_transit = float(overrides.get("in_transit", {}).get(code, transit_totals.get((supplier, code), 0.0)))
        next_arrival = arrival_lookup.get((supplier, code), None)
        if isinstance(next_arrival, pd.Timestamp):
            next_arrival = next_arrival.date()

        if supplier == "IEK":
            abc_class = iek_abc.get(code, "C")
        else:
            abc_class = {"1": "A", "2": "B", "3": "C", "5": "B"}.get(_text(product.manager_class), "C")
        abc_class = overrides.get("abc_class", {}).get(code, abc_class)
        lead = int(settings.lead_time_days[supplier])
        if supplier == "IEK" and _text(product.status).casefold() == "заказная":
            lead = math.ceil(lead * 1.5)
        window_days = lead + int(settings.review_days[supplier])
        window_demand = _window_forecast(settings.today, window_days, base, factors, trend, settings.plan_growth_pct)
        safety_stock = float(settings.service_z[abc_class] * sigma * math.sqrt(window_days / 30))
        raw_qty = max(0.0, window_demand + safety_stock - available - in_transit)
        multiple = int(max(1, moq.get((supplier, code), getattr(product, "catalog_multiple", 1))))
        name_upper = _text(product.name).upper().replace(" ", "")
        if supplier == "IEK" and ("БУХТ" in name_upper or "305М" in name_upper):
            multiple = 305
        qty = int(math.ceil(raw_qty / multiple) * multiple) if raw_qty > 0 else 0

        do_not_reorder = False
        stop_reason = ""
        last6_sales = cleaned_values.reindex(base_months, fill_value=0.0).sum()
        if supplier == "Systeme Electric" and (_text(product.manager_class) == "7" or "!!!" in _text(product.name)):
            do_not_reorder, stop_reason = True, "не продаётся / помечен !!! — вероятно выводится"
        elif supplier == "IEK" and _text(product.status).casefold() == "распродажа":
            do_not_reorder, stop_reason = True, "IEK выводит товар (распродажа)"
        elif last6_sales <= 0 and not any(row["month"] in base_months for row in stockout_rows):
            do_not_reorder, stop_reason = True, "нет продаж 6 мес"
        if do_not_reorder:
            qty = 0

        daily = window_demand / window_days if window_days else 0.0
        enough_until = settings.today + timedelta(days=int((available + in_transit) / daily)) if daily > 0 else None
        if daily > 0 and available + in_transit < daily * lead:
            urgency = "🔴 Критично"
        elif qty > 0:
            urgency = "🟡 Заказать"
        else:
            urgency = "🟢 Не требуется"
        price = float(product.price or 0)
        value = float(qty * price)

        one_offs = one_off_frames.get((supplier, code), invoices.iloc[0:0])
        one_off_details = one_offs[["date", "doc", "qty", "typical_qty"]].reset_index(drop=True)
        unit = _text(product.unit) or "шт"
        reason_parts = [
            f"Регулярный спрос ≈ {base:,.0f} {unit}/мес (мар–авг 2026, без лучшего и худшего месяца).".replace(",", " ")
        ]
        if settings.demand_source == "invoices" and not one_off_details.empty:
            item = one_off_details.sort_values("qty", ascending=False).iloc[0]
            reason_parts.append(
                f"Исключён разовый заказ {item.qty:,.0f} {unit} (накл. {item.doc} от {item.date:%d.%m.%Y}, обычно {item.typical_qty:,.0f} {unit}).".replace(",", " ")
            )
        if stockout_rows:
            item = max(stockout_rows, key=lambda row: row["lost"])
            correction = "поправка отключена" if overrides.get("disable_stockout_adjustment", False) else f"учтено {item['expected']:.0f}"
            reason_parts.append(
                f"В {item['month'].strftime('%m.%Y')} был дефицит: продано {item['actual']:.0f} вместо ~{item['expected']:.0f} — {correction}."
            )
        level_ru = {"product": "товара", "category": "категории", "supplier": "поставщика"}[season_level]
        reason_parts.append(f"Сезонность: текущий месяц ×{factors[settings.today.month - 1]:.2f} (уровень {level_ru}).")
        if abs(trend) >= 0.005:
            reason_parts.append(f"Тренд {trend:+.0%} к прошлому году.")
        if do_not_reorder:
            reason_parts.append(stop_reason.capitalize() + ".")
        elif qty == 0:
            reason_parts.append(f"Запаса хватит до {_reason_date(enough_until)} — заказ не нужен.")
        else:
            reason_parts.append(
                f"Нужно на {window_days} дн. (поставка {lead} + до следующего заказа {settings.review_days[supplier]}): "
                f"{window_demand:.0f} {unit} + страховой запас {safety_stock:.0f} {unit}. Есть {available:.0f} {unit}, в пути {in_transit:.0f} {unit}."
            )
            if in_transit > 0 and next_arrival is not None:
                reason_parts.append(f"Ближайшее поступление: {_reason_date(next_arrival)}.")
            if multiple > 1:
                reason_parts.append(f"Округлено до кратности {multiple}.")
            if multiple == 305 and qty > 0:
                reason_parts.append(f"{qty // 305} бухт ({qty // 305}×305 м).")
        if supplier == "IEK" and _text(product.status).casefold() == "новинка" and (cleaned_values > 0).sum() < 3:
            reason_parts.append("Новинка — проверьте вручную.")
        september_one_off = one_offs[one_offs["date"].dt.to_period("M").eq(pd.Period("2026-09", freq="M"))]
        if supplier == "IEK" and not september_one_off.empty:
            reason_parts.append("Разовый заказ был в сентябре: если отгружено со склада — проверьте остаток в 1С.")

        future_months = pd.period_range(pd.Period(settings.today, freq="M"), periods=12, freq="M")
        forecasts = []
        middle = pd.Period("2026-06", freq="M")
        for month in future_months:
            t = month.ordinal - middle.ordinal
            value_forecast = base * factors[month.month - 1] * (1 + trend) ** (t / 12) * (1 + settings.plan_growth_pct / 100)
            forecasts.append(max(0.0, value_forecast))
        history = pd.DataFrame({
            "raw": raw_values, "cleaned": cleaned_values, "adjusted": adjusted_values,
            "stock_start": stock_values,
        })
        stockout_details = pd.DataFrame(stockout_rows, columns=["month", "actual", "expected", "lost"])
        details[code] = {
            "history": history,
            "season_factors": factors.tolist(),
            "season_level": season_level,
            "forecast": pd.DataFrame({"month": future_months, "forecast": forecasts}),
            "one_offs": one_off_details,
            "stockouts": stockout_details,
        }
        plan_rows.append({
            "supplier": supplier, "code": code, "article": _text(product.article), "name": _text(product.name),
            "category": category, "abc_class": abc_class, "unit": _text(product.unit) or "шт",
            "available": available, "in_transit": in_transit, "next_arrival": next_arrival,
            "base_monthly": base, "trend": trend, "window_days": window_days,
            "window_demand": window_demand, "safety_stock": safety_stock, "raw_qty": raw_qty,
            "multiple": multiple, "qty": qty, "urgency": urgency, "enough_until": enough_until,
            "price": price, "value": value, "do_not_reorder": do_not_reorder,
            "reason": " ".join(reason_parts),
        })

    plan = pd.DataFrame(plan_rows, columns=PLAN_COLUMNS)
    urgency_order = pd.Categorical(plan["urgency"], ["🔴 Критично", "🟡 Заказать", "🟢 Не требуется"], ordered=True)
    plan = plan.assign(_urgency=urgency_order, _risk=np.where(plan["window_days"].gt(0), plan["window_demand"] / plan["window_days"] * plan["price"], 0))
    plan = plan.sort_values(["supplier", "_urgency", "_risk"], ascending=[True, True, False]).drop(columns=["_urgency", "_risk"]).reset_index(drop=True)
    return plan, details
