"""Fill a copy of the SE manager template; never save over the source."""
from copy import copy
from io import BytesIO
import math

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.workbook.properties import CalcProperties


def fill_manager_sheet(template_path: str, plan_df) -> bytes:
    """The caller supplies approved rows only, as specified by app.py."""
    required = {'supplier', 'code', 'qty', 'reason'}
    if not required.issubset(plan_df.columns):
        raise ValueError('Для экспорта нужны supplier, code, qty, reason')
    selected = plan_df.loc[plan_df['supplier'].eq('Systeme Electric')].copy()
    selected['code'] = selected['code'].astype(str).str.strip()
    if selected['code'].duplicated().any():
        raise ValueError('Повтор кода 1С в утверждённом заказе SE')
    lookup = {}
    for row in selected.itertuples():
        qty = float(row.qty)
        if not math.isfinite(qty) or qty < 0 or not qty.is_integer():
            raise ValueError('Количество заказа должно быть целым и неотрицательным')
        lookup[row.code] = (int(qty), str(row.reason))
    workbook = load_workbook(template_path, data_only=False)
    sheets = [ws for ws in workbook.worksheets
              if str(ws['C2'].value).strip().lower() == 'код 1с'
              and str(ws['BB2'].value).strip() == 'Заказ']
    if len(sheets) != 1:
        raise ValueError('Не найден единственный лист с Код 1с в C2 и Заказ в BB2')
    sheet = sheets[0]
    reason_col = sheet.max_column + 1
    sheet.cell(2, reason_col, 'Почему')._style = copy(sheet['BB2']._style)
    sheet.column_dimensions[sheet.cell(2, reason_col).column_letter].width = 70
    # A hidden column group can cover BB, so split it without exposing its neighbours.
    for key, dim in list(sheet.column_dimensions.items()):
        if dim.hidden and dim.min and dim.min <= 54 <= (dim.max or dim.min):
            del sheet.column_dimensions[key]
            for start, end in [(dim.min, 53), (55, dim.max or dim.min)]:
                if start <= end:
                    from openpyxl.utils import get_column_letter
                    part = copy(dim)
                    part.min, part.max = start, end
                    part.index = get_column_letter(start)
                    sheet.column_dimensions[part.index] = part
    sheet.column_dimensions['BB'].hidden = False
    sheet.column_dimensions['BB'].width = max(12, sheet.column_dimensions['BB'].width or 0)
    for row in range(3, sheet.max_row + 1):
        code = str(sheet.cell(row, 3).value or '').strip()
        if not code:
            continue
        qty, reason = lookup.get(code, (0, 'Не включён в утверждённый заказ'))
        sheet.cell(row, 54, qty)
        cell = sheet.cell(row, reason_col, reason)
        cell.data_type = 's'  # Explanations are text, never Excel formulas.
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    workbook.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
