"""Focused checks for Person 2's UI, fixture, KPIs and formula-preserving export."""
from datetime import date
from io import BytesIO
from pathlib import Path
import hashlib
import sys
import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import engine_stub
from export_manager_sheet import fill_manager_sheet
from tab_analytics import _kpis


def test_stub_contract_and_overrides():
    data = engine_stub.load_all()
    settings = engine_stub.Settings()
    plan, details = engine_stub.build_plan(data, settings)
    assert len(plan.columns) == 24
    assert plan['code'].map(type).eq(str).all()
    row = plan.iloc[0]
    after,_ = engine_stub.build_plan(data,settings,{'in_transit':{row.code:100000}})
    assert after.set_index('code').loc[row.code,'qty'] == 0
    for detail in details.values():
        assert isinstance(detail['history'].index,pd.PeriodIndex)
        assert len(detail['forecast']) == 12
        assert sum(detail['season_factors'])/12 == pytest.approx(1.)


def test_preview_and_all_checks():
    app = AppTest.from_file(str(ROOT/'app_preview.py'), default_timeout=60).run()
    assert not app.exception
    assert [tab.label for tab in app.tabs] == ['✅ Проверка','📊 Аналитика']
    app.button(key='checks_all').click().run(timeout=60)
    assert not app.exception
    assert len(app.success) == 5, [x.value for x in app.error]
    assert not app.error


def test_analytics_kpis():
    data = engine_stub.load_all()
    plan, details = engine_stub.build_plan(data,engine_stub.Settings())
    # Set a known stock surplus and ensure lost sales are priced, not counted twice.
    plan['available'] = 1000.
    excess,lost,count = _kpis(plan,details,date(2026,9,22))
    expected = sum(max(0,1000-details[r.code]['forecast']['forecast'].head(3).sum())*r.price for r in plan.itertuples())
    assert excess == pytest.approx(expected)
    assert lost == pytest.approx(4*80*1200)
    assert count == 0


def test_manager_export_preserves_formulas_and_source():
    source = ROOT/'data/se_manager_sheet.xlsx'
    before_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    original = load_workbook(source,data_only=False)
    sheet = next(ws for ws in original if str(ws['C2'].value).strip().lower() == 'код 1с')
    row_number = next(r for r in range(3,sheet.max_row+1) if sheet.cell(r,3).value)
    code = str(sheet.cell(row_number,3).value).strip()
    plan = pd.DataFrame([{'supplier':'Systeme Electric','code':code,'qty':25,'reason':'=Текст, не формула'},
                         {'supplier':'IEK','code':code,'qty':999,'reason':'Другой поставщик'}])
    output = load_workbook(BytesIO(fill_manager_sheet(str(source),plan)),data_only=False)
    changed = output[sheet.title]
    assert changed.cell(row_number,54).value == 25
    assert changed.cell(row_number,sheet.max_column+1).value == '=Текст, не формула'
    assert changed.cell(row_number,sheet.max_column+1).data_type == 's'
    assert changed.cell(2,sheet.max_column+1).value == 'Почему'
    assert not changed.column_dimensions['BB'].hidden
    assert output.calculation.fullCalcOnLoad
    for ws in original:
        for row in ws:
            for cell in row:
                if ws.title == sheet.title and cell.column == 54 and cell.row >= 3:
                    continue
                saved = output[ws.title][cell.coordinate].value
                if isinstance(cell.value, ArrayFormula):
                    assert isinstance(saved, ArrayFormula)
                    assert (saved.text, saved.ref) == (cell.value.text, cell.value.ref)
                else:
                    assert saved == cell.value
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before_hash
    blank = load_workbook(BytesIO(fill_manager_sheet(str(source),plan.iloc[:0])))
    assert blank[sheet.title].cell(row_number,54).value == 0
    original.close();output.close();blank.close()


@pytest.mark.parametrize('quantity', [-1,1.5,float('nan'),float('inf')])
def test_export_rejects_invalid_quantities(quantity):
    plan = pd.DataFrame([{'supplier':'Systeme Electric','code':'x','qty':quantity,'reason':'Причина'}])
    with pytest.raises(ValueError):
        fill_manager_sheet(str(ROOT/'data/se_manager_sheet.xlsx'),plan)
